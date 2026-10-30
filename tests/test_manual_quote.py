"""Manual editor data, backwards compatibility and export parity."""
import io,json,math,sys,unittest,zipfile
import test_saving
from manual_quote import adjusted_cost
from pdf_client_quote import summarize_quote,build_client_quote
from excel_quote import build_quote
from openpyxl import load_workbook

class ManualQuoteTest(unittest.TestCase):
    setUp=test_saving.SavingTest.setUp
    tearDown=test_saving.SavingTest.tearDown
    start_server=test_saving.SavingTest.start_server
    stop_server=test_saving.SavingTest.stop_server
    request=test_saving.SavingTest.request
    def test_manual_roundtrip_and_library(self):
        row={'descrizione':'Programmazione PLC','quantita':2,'prezzo':100,'unita':'h','sezione':'Programmazione','sconto':10,'ricarico_voce':20}
        payload={'cliente':'Prova','titolo':'Manuale','righe':[row,dict(row,descrizione='Trasferta',sezione='Trasferte')],'note':'Nota cliente','condizioni':'Esclusioni','validita':'30 giorni','pagamenti':'Bonifico','riferimento':'OFF-1','iva':22}
        status,data=self.request('/api/preventivi',payload)
        self.assertEqual(status,200,data)
        pid=data['id']
        status,saved=self.request(f'/api/preventivi/{pid}')
        self.assertEqual(saved['righe'][0]['sezione'],'Programmazione')
        self.assertEqual(saved['righe'][0]['sconto'],10)
        self.assertEqual(json.loads(saved['preventivo']['dettagli_json'])['note'],'Nota cliente')
        payload['id']=pid;payload['righe'].reverse()
        self.assertEqual(self.request('/api/preventivi',payload)[0],200)
        self.assertEqual(self.request(f'/api/preventivi/{pid}')[1]['righe'][0]['sezione'],'Trasferte')
        before=self.request(f'/api/preventivi/{pid}')[1]
        payload['righe'][0]['sconto']=101
        self.assertEqual(self.request('/api/preventivi',payload)[0],400)
        self.assertEqual(self.request(f'/api/preventivi/{pid}')[1],before)
        status,item=self.request('/api/voci-personali',{'voce':row})
        self.assertEqual(status,200,item)
        self.assertEqual(len(self.request('/api/voci-personali')[1]),1)
        self.stop_server();self.scope=test_saving.load_app(self.folder);self.start_server()
        self.assertEqual(self.request(f'/api/preventivi/{pid}')[1],before)
        self.request('/api/voci-personali',{'id':item['id'],'elimina':True})
        self.assertEqual(self.request('/api/voci-personali')[1],[])

    def test_legacy_migration_preserves_original_values(self):
        status,data=self.request('/api/preventivi',{'cliente':'Vecchio','titolo':'Archivio','ricarico':15,'iva':22,'righe':[{'codice':'OLD','descrizione':'Storico','prezzo':37.15,'quantita':3,'quota_materiale':100}]})
        self.assertEqual(status,200)
        self.stop_server()
        c=self.scope['conn']()
        for table,columns in [('preventivi',['dettagli_json']),('righe',['sezione','sconto','ricarico_voce'])]:
            for column in columns:c.execute('ALTER TABLE '+table+' DROP COLUMN '+column)
        c.commit()
        originals={table:[dict(r) for r in c.execute('SELECT * FROM '+table)] for table in ('preventivi','righe')}
        c.close()
        self.scope=test_saving.load_app(self.folder);self.start_server()
        c=self.scope['conn']()
        for table,rows in originals.items():
            actual=[dict(r) for r in c.execute('SELECT * FROM '+table)]
            for old,new in zip(rows,actual):self.assertEqual(old,{k:new[k] for k in old})
            self.assertEqual(len(rows),len(actual))
        self.assertEqual(c.execute('PRAGMA integrity_check').fetchone()[0],'ok');c.close()

class ManualCalculationTest(unittest.TestCase):
    def test_exports(self):
        r={'descrizione':'PLC','sezione':'Programmazione','unita':'h','quantita':2,'prezzo':100,'sconto':10,'ricarico_voce':20,'quota_materiale':50}
        q={'cliente':{'nome':'Prova'},'righe':[r],'ricarico':10,'sconto_materiale':20,'iva':22,'note':'Nota cliente','pagamenti':'Bonifico'}
        expected=2*100*.9*1.2*.9*1.1
        a,vat=summarize_quote(q)
        self.assertAlmostEqual(a[0]['totale'],expected)
        wb=load_workbook(io.BytesIO(build_quote(q)),data_only=True)
        self.assertAlmostEqual(wb['Calcolo interno']['J9'].value,expected)
        self.assertAlmostEqual(wb['Riepilogo generale']['E19'].value,expected)
        self.assertIn('Condizioni commerciali',wb.sheetnames)
        self.assertTrue(build_client_quote(q,'ghiro_preventivi/logo.jpg').startswith(b'%PDF-'))
    def test_old_and_invalid(self):
        self.assertEqual(adjusted_cost({'prezzo':100}),100)
        for value in [-1,101,float('nan'),float('inf')]:
            with self.assertRaises(ValueError):adjusted_cost({'prezzo':100,'sconto':value})
